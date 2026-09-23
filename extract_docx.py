#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Extrair informações do arquivo DOCX de Maturidade Mobilize."""

from docx import Document
import json
import os

# Caminhos
docx_path = r"C:\Users\User\Downloads\D.C.857.GOCO. Analise de Maturidade Mobilize.docx"

def extract_docx_content(docx_path):
    """Extrair todo o conteúdo do arquivo DOCX."""
    doc = Document(docx_path)
    
    content = {
        "paragraphs": [],
        "tables": [],
        "full_text": []
    }
    
    # Extrair parágrafos
    for para in doc.paragraphs:
        if para.text.strip():
            content["paragraphs"].append({
                "text": para.text,
                "style": para.style.name if para.style else None
            })
            content["full_text"].append(para.text)
    
    # Extrair tabelas
    for i, table in enumerate(doc.tables):
        table_data = []
        for row in table.rows:
            row_data = []
            for cell in row.cells:
                row_data.append(cell.text.strip())
            table_data.append(row_data)
        content["tables"].append({
            "index": i,
            "rows": len(table_data),
            "cols": len(table_data[0]) if table_data else 0,
            "data": table_data
        })
    
    return content

if __name__ == "__main__":
    print("Extraindo conteúdo do arquivo DOCX...")
    content = extract_docx_content(docx_path)
    
    # Salvar conteúdo em JSON para análise
    output_path = r"C:\Users\User\maturidade-dashboard\mobilize_content.json"
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(content, f, ensure_ascii=False, indent=2)
    
    print(f"Conteúdo extraído e salvo em: {output_path}")
    print(f"Total de parágrafos: {len(content['paragraphs'])}")
    print(f"Total de tabelas: {len(content['tables'])}")
    
    # Mostrar resumo
    print("\n=== RESUMO DO DOCUMENTO ===")
    for i, table in enumerate(content["tables"]):
        print(f"\nTabela {i+1}: {table['rows']} linhas x {table['cols']} colunas")
        if table['data']:
            print("Primeiras 3 linhas:")
            for row in table['data'][:3]:
                print(f"  {row}")